from datetime import date
from io import BytesIO

from openpyxl import Workbook

from fxlab.communications import parse_ea_ced


def workbook_bytes():
    book = Workbook()
    book.remove(book.active)
    events = book.create_sheet("1. IMC Events (unfiltered)")
    events.append(["ID", "Date", "hour", "minute", "t_account", "s_executiveboard", "s_president",
                   "s_ECBpreshearing", "s_BdF", "s_Buba", "s_BdI", "s_BdE", "interview",
                   "t_speechinECBDB", "ECBDB_ID", "ECBDB_ID2", "ECBDB_ID3", "ECBDB_ID4"])
    events.append([7, date(2020, 1, 15), 14, 30, 0, 0, 1, 0, 0, 0, 0, 0, 0, 1, 42, None, None, None])
    returns = book.create_sheet("2. IMC All returns (filtered)")
    returns.append(["ID", "Date", "Hour", "Minute", "EUROSTOXX", "OIS_2Y", "OIS_10Y", "EURUSD"])
    returns.append([7, "15-Jan-2020", 14, 30, 0.2, 1.5, 2.0, 0.25])
    abnormal = book.create_sheet("3. IMC Abnormal returns (filt) ")
    abnormal.append(["ID", "EURUSD"])
    abnormal.append([7, 0.25])
    speeches = book.create_sheet("6. ECB speeches database")
    speeches.append(["date", "speakers", "title", "subtitle", "contents", "ECBDBID"])
    speeches.append([date(2020, 1, 15), "Speaker", "Title", "Subtitle", "Full speech text", 42])
    target = BytesIO()
    book.save(target)
    return target.getvalue()


def test_parse_ea_ced_links_text_and_market_reaction():
    events, examples, summary = parse_ea_ced(workbook_bytes(), date(2020, 1, 1), date(2020, 12, 31))
    assert len(events) == len(examples) == 1
    assert events[0]["event_type"] == "ecb_president_communication"
    assert events[0]["event_time"].isoformat() == "2020-01-15T13:30:00+00:00"
    assert events[0]["eurusd_return_bp"] == 25.0
    assert events[0]["eurusd_abnormal"] is True
    assert examples[0]["text"] == "Full speech text"
    assert examples[0]["strict_pit_eligible"] is False
    assert summary["median_absolute_eurusd_return_pct"] == 0.25
