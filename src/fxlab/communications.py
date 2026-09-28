"""ECB communication events, texts and intraday market reactions from EA-CED."""

import hashlib
import io
import json
from datetime import date, datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import httpx
from openpyxl import load_workbook

from .macro import write_macro_parquet
from .store import atomic_json, root, save_raw

UTC = timezone.utc
FRANKFURT = ZoneInfo("Europe/Berlin")
SOURCE_URL = (
    "https://www.dropbox.com/scl/fi/s2z0feb6xauj0w9h3ebx9/"
    "EA-CED-Database_IOS.xlsx?rlkey=al3usdrx6h7nuxr5ecu9s70v1&dl=1"
)
SOURCE_PAGE = "https://sites.google.com/site/istrefiklodiana/ea-ced"
PARSER = "ea-ced-2024-07-18-v1"
EVENT_FLAGS = {
    "t_account": "ecb_monetary_policy_account",
    "s_president": "ecb_president_communication",
    "s_ECBpreshearing": "ecb_president_hearing",
    "s_executiveboard": "ecb_executive_board_communication",
    "s_BdF": "banque_de_france_governor_communication",
    "s_Buba": "bundesbank_president_communication",
    "s_BdI": "banca_ditalia_governor_communication",
    "s_BdE": "banco_de_espana_governor_communication",
}


def _sheet_rows(workbook, name: str) -> list[dict]:
    iterator = workbook[name].iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(iterator)]
    return [dict(zip(headers, row)) for row in iterator if any(value is not None for value in row)]


def _integer(value):
    return int(value) if value is not None else None


def _event_kind(row: dict) -> str:
    kinds = [label for flag, label in EVENT_FLAGS.items() if row.get(flag) == 1]
    return "+".join(kinds) if kinds else "other_eurosystem_communication"


def _event_timestamp(row: dict) -> datetime:
    day = row["Date"]
    if not isinstance(day, datetime):
        raise ValueError("EA-CED event date is not an Excel datetime")
    local = datetime(day.year, day.month, day.day, _integer(row["hour"]), _integer(row["minute"]), tzinfo=FRANKFURT)
    return local.astimezone(UTC)


def parse_ea_ced(body: bytes, start: date, end: date) -> tuple[list[dict], list[dict], dict]:
    if not body.startswith(b"PK"):
        raise ValueError("EA-CED response is not an XLSX workbook")
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    required = {"1. IMC Events (unfiltered)", "2. IMC All returns (filtered)",
                "3. IMC Abnormal returns (filt) ", "6. ECB speeches database"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError("EA-CED workbook layout changed")

    events = _sheet_rows(workbook, "1. IMC Events (unfiltered)")
    returns = _sheet_rows(workbook, "2. IMC All returns (filtered)")
    abnormal = _sheet_rows(workbook, "3. IMC Abnormal returns (filt) ")
    speeches = _sheet_rows(workbook, "6. ECB speeches database")
    event_by_id = {_integer(row["ID"]): row for row in events}
    abnormal_fx = {_integer(row["ID"]): float(row["EURUSD"]) for row in abnormal if row.get("EURUSD") is not None}
    speech_by_id = {_integer(row["ECBDBID"]): row for row in speeches if row.get("ECBDBID") is not None}

    normalized, text_examples = [], []
    missing_event_ids = 0
    for market in returns:
        event_id = _integer(market.get("ID"))
        event = event_by_id.get(event_id)
        if event is None:
            missing_event_ids += 1
            continue
        timestamp = _event_timestamp(event)
        if not start <= timestamp.date() <= end:
            continue
        fx_pct = float(market["EURUSD"]) if market.get("EURUSD") is not None else None
        speech_ids = [_integer(event.get(column)) for column in ("ECBDB_ID", "ECBDB_ID2", "ECBDB_ID3", "ECBDB_ID4")]
        speech_rows = [speech_by_id[item] for item in speech_ids if item in speech_by_id]
        full_text = "\n\n".join(str(item.get("contents") or "").strip() for item in speech_rows).strip() or None
        record = {
            "event_id": f"ea-ced-{event_id}", "source_event_id": event_id, "source": "ea_ced",
            "event_type": _event_kind(event), "event_time": timestamp, "published_at": None,
            "available_at": None, "time_quality": "exact_timestamp_reported_by_bloomberg",
            "strict_pit_eligible": False, "interview": event.get("interview") == 1,
            "speech_in_ecb_database": event.get("t_speechinECBDB") == 1,
            "ecb_speech_ids": ",".join(str(item) for item in speech_ids if item is not None) or None,
            "speaker": " | ".join(str(item.get("speakers") or "").strip() for item in speech_rows) or None,
            "title": " | ".join(str(item.get("title") or "").strip() for item in speech_rows) or None,
            "subtitle": " | ".join(str(item.get("subtitle") or "").strip() for item in speech_rows) or None,
            "full_text": full_text, "eurusd_return_pct": fx_pct,
            "eurusd_return_bp": fx_pct * 100 if fx_pct is not None else None,
            "eurusd_abnormal": event_id in abnormal_fx,
            "eurusd_abnormal_return_pct": abnormal_fx.get(event_id),
            "eurostoxx_return_pct": float(market["EUROSTOXX"]) if market.get("EUROSTOXX") is not None else None,
            "ois_2y_change_bp": float(market["OIS_2Y"]) if market.get("OIS_2Y") is not None else None,
            "ois_10y_change_bp": float(market["OIS_10Y"]) if market.get("OIS_10Y") is not None else None,
            "source_url": SOURCE_PAGE,
        }
        normalized.append(record)
        if full_text and fx_pct is not None:
            text_examples.append({
                "example_id": record["event_id"], "event_time": timestamp, "speaker": record["speaker"],
                "title": record["title"], "text": full_text, "eurusd_return_pct": fx_pct,
                "eurusd_return_bp": fx_pct * 100, "eurusd_abnormal": record["eurusd_abnormal"],
                "use": "retrospective_text_feature_research",
                "strict_pit_eligible": False,
                "pit_warning": "Full speech availability time is unknown; do not use as a pre-event prediction input.",
            })
    normalized.sort(key=lambda row: (row["event_time"], row["event_id"]))
    text_examples.sort(key=lambda row: (row["event_time"], row["example_id"]))
    if not normalized:
        raise ValueError("EA-CED replay produced no rows")
    fx = [row["eurusd_return_pct"] for row in normalized if row["eurusd_return_pct"] is not None]
    abnormal_values = [row["eurusd_abnormal_return_pct"] for row in normalized if row["eurusd_abnormal_return_pct"] is not None]
    def median_absolute(values):
        ordered = sorted(abs(value) for value in values)
        middle = len(ordered) // 2
        return ordered[middle] if len(ordered) % 2 else (ordered[middle - 1] + ordered[middle]) / 2
    summary = {
        "workbook_events": len(events), "filtered_events": len(returns), "normalized_events": len(normalized),
        "missing_event_ids": missing_event_ids, "speech_text_examples": len(text_examples),
        "eurusd_observations": len(fx), "eurusd_abnormal_observations": len(abnormal_values),
        "median_absolute_eurusd_return_pct": median_absolute(fx),
        "mean_absolute_eurusd_return_pct": sum(abs(value) for value in fx) / len(fx),
        "median_absolute_abnormal_eurusd_return_pct": median_absolute(abnormal_values),
    }
    return normalized, text_examples, summary


def fetch_ea_ced() -> dict:
    with httpx.Client(timeout=180, follow_redirects=True, headers={"User-Agent": "FXCausalLab/0.28 research"}) as client:
        response = client.get(SOURCE_URL)
        meta = save_raw("ea_ced", str(response.url), response.content, dict(response.headers), response.status_code)
        response.raise_for_status()
    manifest = {"provider": "ea_ced", "snapshot": meta, "fetched_at": datetime.now(UTC).isoformat()}
    atomic_json(root()/"reports"/"ea_ced_fetch.json", manifest)
    return manifest


def replay_ea_ced(manifest: Path | dict, start: date = date(1999, 1, 1), end: date = date.max) -> dict:
    source = json.loads(manifest.read_text(encoding="utf-8")) if isinstance(manifest, Path) else manifest
    meta = source["snapshot"]
    body = (root()/meta["payload"]).read_bytes()
    if hashlib.sha256(body).hexdigest() != meta["sha256"]:
        raise ValueError("EA-CED raw checksum mismatch")
    events, examples, summary = parse_ea_ced(body, start, end)
    signature = {"parser": PARSER, "raw_sha256": meta["sha256"], "start": str(start), "end": str(end)}
    dataset_id = hashlib.sha256(json.dumps(signature, sort_keys=True).encode()).hexdigest()[:20]
    folder = root()/"silver"/"communications"/dataset_id
    write_macro_parquet(events, folder/"ea_ced_events.parquet", ("event_time", "published_at", "available_at"))
    write_macro_parquet(examples, folder/"anyjev_text_examples.parquet", ("event_time",))
    report = {
        "dataset_id": dataset_id, "provider": "ea_ced", "parser": PARSER,
        "requested_start": str(start), "requested_end": str(end), **summary,
        "first_event_time": events[0]["event_time"].isoformat(), "last_event_time": events[-1]["event_time"].isoformat(),
        "files": {"events": (folder/"ea_ced_events.parquet").relative_to(root()).as_posix(),
                  "anyjev_examples": (folder/"anyjev_text_examples.parquet").relative_to(root()).as_posix()},
        "snapshot": meta, "strict_pit_eligible": False,
        "limitations": [
            "Event hour and minute are reported from Bloomberg, but historical system receipt is not proven.",
            "Full ECB speech text has no exact publication timestamp and cannot be used as a pre-event input.",
            "The supplied market reaction is useful for event research; future 1/5/20/60-day targets must be joined separately.",
            "Events near major macro releases and FOMC decisions are already excluded from the filtered reaction sheet.",
        ], "generated_at": datetime.now(UTC).isoformat(),
    }
    atomic_json(folder/"manifest.json", report)
    atomic_json(root()/"reports"/"communications.json", report)
    return report


def backfill_ea_ced(start: date = date(1999, 1, 1), end: date = date.max) -> dict:
    return replay_ea_ced(fetch_ea_ced(), start, end)
